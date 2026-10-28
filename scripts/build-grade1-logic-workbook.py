"""Build the original illustrated workbook from the reviewed JSON content.
Run with the Codex bundled Python. Generated files stay in ignored outputs/.
"""
from pathlib import Path
import argparse, json, math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
CONTENT=ROOT/'docs/grade1-logic-content.json'
INK='#253e47'; TEAL='#4b9e93'; GOLD='#efb95c'; PINK='#eab4a3'; BLUE='#a5c7d9'

def build(font_path, out):
 data=json.loads(CONTENT.read_text());out.mkdir(parents=True,exist_ok=True);assets=out/'assets';assets.mkdir(exist_ok=True)
 def font(n):return ImageFont.truetype(str(font_path),n)
 def text(draw,xy,t,n=32,fill=INK,anchor=None):draw.text(xy,t,font=font(n),fill=fill,anchor=anchor)
 def arrow(draw,a,b,fill=TEAL,w=6):
  draw.line([a,b],fill=fill,width=w);ang=math.atan2(b[1]-a[1],b[0]-a[0]);tip=[b,(b[0]-22*math.cos(ang-.5),b[1]-22*math.sin(ang-.5)),(b[0]-22*math.cos(ang+.5),b[1]-22*math.sin(ang+.5))];draw.polygon(tip,fill=fill)
 def bunny(draw,cx,cy,s=1):
  def box(a,b,c,d):return (cx+a*s,cy+b*s,cx+c*s,cy+d*s)
  draw.ellipse(box(-62,-140,-13,-8),fill='#fff4e6',outline=INK,width=3);draw.ellipse(box(12,-140,61,-8),fill='#fff4e6',outline=INK,width=3)
  draw.ellipse(box(-50,-125,-25,-28),fill=PINK);draw.ellipse(box(25,-125,48,-28),fill=PINK)
  draw.ellipse(box(-82,-65,82,78),fill='#fff4e6',outline=INK,width=3)
  draw.ellipse(box(-37,-15,-25,0),fill=INK);draw.ellipse(box(25,-15,37,0),fill=INK);draw.ellipse(box(-8,10,8,22),fill=PINK)
  draw.arc(box(-23,10,0,40),0,90,fill=INK,width=3);draw.arc(box(0,10,23,40),90,180,fill=INK,width=3)
  draw.ellipse(box(-58,8,-32,23),fill=PINK);draw.ellipse(box(32,8,58,23),fill=PINK)
  draw.rounded_rectangle(box(-57,70,57,150),radius=18,fill=TEAL,outline=INK,width=3)
 def fox(draw,cx,cy,s=1):
  pts=[(cx+x*s,cy+y*s) for x,y in [(-82,-85),(-70,10),(0,87),(70,10),(82,-85),(35,-58),(-35,-58)]];draw.polygon(pts,fill=GOLD,outline=INK,width=3)
  draw.polygon([(cx-65*s,cy+5*s),(cx,cy+65*s),(cx+65*s,cy+5*s),(cx+30*s,cy+65*s),(cx-30*s,cy+65*s)],fill='#fff4e6')
  for x in [-30,30]:draw.ellipse((cx+(x-5)*s,cy-12*s,cx+(x+5)*s,cy+2*s),fill=INK)
  draw.ellipse((cx-9*s,cy+35*s,cx+9*s,cy+48*s),fill=INK)
  draw.rounded_rectangle((cx-52*s,cy+78*s,cx+52*s,cy+146*s),radius=15,fill=BLUE,outline=INK,width=3)
 def scene(i):
  im=Image.new('RGB',(1400,420),'white');d=ImageDraw.Draw(im)
  d.ellipse((-100,275,1500,620),fill='#eff7ee');d.ellipse((1120,15,1210,105),fill='#f7d483')
  for x in [50,1300]:
   d.rectangle((x,155,x+14,290),fill='#b89170');d.ellipse((x-48,80,x+64,215),fill='#a7cbb0')
  bunny(d,215,225,.8);fox(d,1175,225,.8)
  themes=['房间卡去哪儿了','一起数灯间的小路','给火车选符号','小店长请就位','把故事倒着走','找出图形里的数字','先猜一猜再展开']
  text(d,(700,65),themes[i],46,anchor='mm');text(d,(215,385),'米米',28,anchor='mm');text(d,(1175,385),'点点',28,anchor='mm')
  if i==0:
   for y,row in enumerate([[1,'?',3],[2,3,1],[3,1,2]]):
    for x,v in enumerate(row):
     xx=560+x*90;yy=140+y*75;d.rounded_rectangle((xx,yy,xx+80,yy+65),radius=8,fill='#eef6f5',outline=TEAL,width=3);text(d,(xx+40,yy+32),str(v),38,anchor='mm')
  elif i==1:
   for x in [420,550,680,810,940]:
    d.line((x,245,x,345),fill=INK,width=5);d.ellipse((x-20,210,x+20,250),fill=GOLD,outline=INK,width=3)
   d.line((420,345,940,345),fill=TEAL,width=4)
  elif i==2:
   for j,v in enumerate(['7','?','2','?','3','=8']):
    xx=400+j*100;d.rounded_rectangle((xx,190,xx+85,295),radius=12,fill=BLUE if j%2==0 else '#f8e7c7',outline=INK,width=3);text(d,(xx+42,240),v,38,anchor='mm')
    for dx in [20,65]:d.ellipse((xx+dx-7,300,xx+dx+7,314),fill=INK)
  elif i==3:
   d.polygon([(420,205),(500,140),(900,140),(980,205)],fill=PINK);d.rectangle((450,205,950,330),fill='#f9f0dd',outline=INK,width=3)
   text(d,(700,250),'森林文具店',45,anchor='mm');text(d,(700,305),'先看价签  再算总价',27,anchor='mm')
  elif i==4:
   for j,v in enumerate(['原来 ?','拿走 3','放入 2','最后 7']):
    xx=390+j*160;d.ellipse((xx,210,xx+125,335),fill=['#f9e8c5','#e8f4ef','#e8f4ef','#e1edf7'][j],outline=INK,width=3);text(d,(xx+62,272),v,27,anchor='mm')
    if j<3:arrow(d,(xx+128,272),(xx+157,272),w=4)
  elif i==5:
   text(d,(700,215),'△ ＋ △ ＝ 8',62,anchor='mm');text(d,(700,300),'△ ＋ ○ ＝ 7',62,anchor='mm')
  else:
   d.polygon([(460,160),(650,160),(650,345),(460,345)],fill=BLUE,outline=INK,width=3);d.line((555,160,555,345),fill=INK,width=3)
   arrow(d,(690,255),(760,255));d.polygon([(850,200),(900,160),(950,200),(950,260),(900,335),(850,260)],fill=PINK,outline=INK,width=3)
  p=assets/f'scene-{i+1}.png';im.save(p);return p
 def diagram(i):
  im=Image.new('RGB',(1400,390),'white');d=ImageDraw.Draw(im)
  if i==0:
   for y,row in enumerate([[1,'?',3],['?',3,1],[3,1,'?']]):
    for x,v in enumerate(row):
     xx=90+x*105;yy=30+y*105;d.rectangle((xx,yy,xx+105,yy+105),outline=TEAL,width=4);text(d,(xx+52,yy+52),str(v),45,anchor='mm')
   arrow(d,(475,80),(630,80));text(d,(675,70),'这一行有 1 和 3，还缺 2。',38)
   arrow(d,(475,185),(630,185));text(d,(675,175),'每行检查后，再竖着检查。',38)
   text(d,(700,300),'规则：每行、每列都有 1、2、3。',31)
  elif i==1:
   for j in range(5):
    x=180+j*250;d.line((x,115,x,220),fill=INK,width=5);d.ellipse((x-27,70,x+27,125),fill=GOLD,outline=INK,width=3);text(d,(x,50),str(j+1),30,anchor='mm')
    if j<4:
     d.line((x+5,220,x+245,220),fill=TEAL,width=8);text(d,(x+125,270),'2 米',36,anchor='mm')
   text(d,(700,340),'5 盏灯之间有 4 段，每段长 2 米。',40,anchor='mm')
  elif i==2:
   for j,(a,b,v) in enumerate([('+','+',12),('+','－',6),('－','+',8),('－','－',2)]):
    x=60+(j%2)*680;y=25+(j//2)*180;d.rounded_rectangle((x,y,x+620,y+145),radius=20,fill='#eef6f5' if v==8 else '#f7f4ed',outline=TEAL if v==8 else '#c7d5d6',width=3);text(d,(x+310,y+52),f'7 {a} 2 {b} 3 = {v}',43,anchor='mm');text(d,(x+310,y+108),'到达 8 号站' if v==8 else '还没有到 8 号站',28,anchor='mm')
  elif i==3:
   for j,(name,price) in enumerate([('铅笔','1 元 5 角'),('橡皮','5 角'),('付出','5 元')]):
    x=95+j*440;d.rounded_rectangle((x,35,x+340,280),radius=20,fill=['#e6f2ed','#fae9dc','#e8f0f7'][j],outline=INK,width=3)
    if j==0:
     d.polygon([(x+100,85),(x+260,85),(x+280,100),(x+260,115),(x+100,115)],fill=GOLD,outline=INK)
    elif j==1:d.rounded_rectangle((x+135,70,x+225,125),radius=8,fill=PINK,outline=INK,width=3)
    else:text(d,(x+170,100),'练习币',34,anchor='mm')
    text(d,(x+170,175),name,40,anchor='mm');text(d,(x+170,235),price,38,anchor='mm')
   text(d,(700,340),'合起来 2 元，付出 5 元，应找回 3 元。',39,anchor='mm')
  elif i==4:
   for j,v in enumerate(['原来 ?','中间 ?','最后 7']):
    x=70+j*460;d.rounded_rectangle((x,125,x+325,245),radius=20,fill='#e9f4ef',outline=TEAL,width=3);text(d,(x+162,185),v,40,anchor='mm')
   for x,t in [(400,'拿走 3'),(860,'放入 2')]:
    arrow(d,(x,150),(x+120,150));text(d,(x+60,85),t,30,anchor='mm');arrow(d,(x+120,240),(x,240),fill='#ba7557');text(d,(x+60,295),'加回 3' if x==400 else '减去 2',30,anchor='mm')
   text(d,(700,360),'上面顺着讲故事，下面倒着找原数。',35,anchor='mm')
  elif i==5:
   text(d,(90,75),'△ ＋ △ ＝ 8',52);text(d,(90,200),'△ ＋ ○ ＝ 7',52);arrow(d,(610,110),(740,110));arrow(d,(610,235),(740,235))
   text(d,(790,75),'4 ＋ 4 ＝ 8',52);text(d,(790,200),'4 ＋ 3 ＝ 7',52);text(d,(700,345),'先求出 △，再替换第二句中的 △。',38,anchor='mm')
  else:
   for j,(label,holes) in enumerate([('对折一次',1),('打开看看',2)]):
    x=120+j*700;d.rectangle((x,70,x+440,295),fill='#e6f1f7',outline=INK,width=4)
    for yy in range(70,295,20):d.line((x+220,yy,x+220,yy+10),fill=TEAL,width=4)
    positions=[330] if holes==1 else [110,330]
    for dx in positions:d.polygon([(x+dx,145),(x+dx-25,195),(x+dx+25,195)],fill='white',outline=INK,width=3)
    if holes==1:
     d.rectangle((x-3,66,x+218,299),fill='white')
     for yy in range(70,295,20):d.line((x+220,yy,x+220,yy+10),fill=TEAL,width=4)
    text(d,(x+220,345),label,36,anchor='mm')
   arrow(d,(640,185),(760,185));text(d,(700,24),'剪口远离折痕与纸的外边',31,anchor='mm')
  p=assets/f'diagram-{i+1}.png';im.save(p);return p
 doc=Document();sec=doc.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21);sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.6);sec.left_margin=sec.right_margin=Cm(1.8)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  st=doc.styles[name];st.font.name='Arial Unicode MS';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS');st.font.color.rgb=RGBColor(0,0,0)
 for e in list(doc.styles.element.iter()):
  if e.tag==qn('w:pBdr'):e.getparent().remove(e)
 doc.styles['Normal'].font.size=Pt(12);doc.styles['Normal'].paragraph_format.line_spacing=1.05;doc.styles['Normal'].paragraph_format.space_after=Pt(5)
 for name,size in [('Title',26),('Heading 1',22),('Heading 2',15)]:
  doc.styles[name].font.size=Pt(size);doc.styles[name].paragraph_format.space_before=Pt(7);doc.styles[name].paragraph_format.space_after=Pt(5)
 footer=sec.footer.paragraphs[0];footer.alignment=2;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 def p(t):return doc.add_paragraph(t)
 def h(t):return doc.add_heading(t,2)
 def page(title):
  a=doc.add_heading(title,1);a.paragraph_format.page_break_before=True
 def pic(path,width=16):
  a=doc.add_paragraph();a.paragraph_format.space_after=Pt(5);r=a.add_run();shape=r.add_picture(str(path),width=Cm(width));shape._inline.docPr.set('descr',path.stem.replace('-',' '))
 def space(lines=1):
  for _ in range(lines):a=p('________________________________________________________________');a.paragraph_format.space_after=Pt(9)
 def grid(rows):
  t=doc.add_table(rows=3,cols=3);t.alignment=WD_TABLE_ALIGNMENT.LEFT;t.autofit=False
  for col in t.columns: col.width=Cm(1.5)
  tw=t._tbl.tblPr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),'2551')
  for y in range(3):
   for x in range(3):
    c=t.cell(y,x);c.width=Cm(1.5);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=str(rows[y][x])
    pp=c.paragraphs[0];pp.alignment=1;pp.paragraph_format.space_before=Pt(4);pp.paragraph_format.space_after=Pt(4)
    cp=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'8');e.set(qn('w:color'),'D9D9D9');b.append(e)
    cp.append(b)
 def question(l,i,n):
  level,q=l['qs'][n-1];h(f'{n}  {level}')
  if i==0 and n in [1,3,5]:
   lead=q.split('\n')[0];p(lead);rows={1:[[1,'',3],['',3,1],[3,1,'']],3:[[1,'',''],['',1,2],[2,'',1]],5:[[1,'',''],['','',''],['','','']]}[n];grid(rows)
  else:p(q)
  space(1 if i==0 else 2)
 doc.add_heading(data['title'],0);p('图解故事版  |  7 个主题  |  亲子共学')
 pic(scene(0));h('给小小探险家')
 p('米米、点点和慢慢正在准备森林游园会。旅馆、花灯、小店、工坊都遇到了小难题。带上你的铅笔，帮他们找线索吧！')
 h('给陪伴孩子的大人');p('每讲分为故事讲解、示范跟练、探索练习三页。建议每次 15 至 20 分钟，分两到三次完成一讲。成人可以读题，孩子负责摆、画、说。挑战题选做，答错后可以再试。')
 p('准备：20 个瓶盖、数字卡、符号卡、练习币、薄纸、彩笔。使用剪刀时，由成人看护。书后集中提供 42 道练习的答案，方便分开打印。')
 h('七次探险');p('① 房间卡  ② 花灯小路  ③ 算式火车\n④ 文具店  ⑤ 果子追踪  ⑥ 图形密码  ⑦ 蝴蝶工坊')
 p('本书为原创配套学习材料，非学而思官方教材。')
 for i,l in enumerate(data['lessons']):
  page(f'第{i+1}讲  {l["title"]}');p(l['scene']);pic(scene(i));p(l['dialogue']);h('今天的任务');p(l['mission']);p('学会：'+l['goal']);h('先玩一玩再发现');p(l['explain']);h('我的思考路线')
  for n,s in enumerate(l['steps'],1):p(f'{n}  {s}')
  page(f'第{i+1}讲  一起解开谜题');pic(diagram(i));h('跟着想一想');p(l['example']);h('换你来');p(l['follow']);space();question(l,i,1);question(l,i,2)
  page(f'第{i+1}讲  我来当小探险家');p('先做进阶题，挑战题可以选做。需要时，用瓶盖或图画帮忙。')
  for n in [3,4,5]:question(l,i,n)
  h('亲子游戏');p(l['game']);p('我能说清：我先 __________，再 __________，最后用 __________ 检查。')
 for start in range(0,7,2):
  page('家长参考答案与带学提示');p('先听孩子解释，再对照答案。每讲“换你来”1 题与分层练习 5 题，共 6 题。')
  for idx in range(start,min(start+2,7)):
   l=data['lessons'][idx];h(f'第{idx+1}讲  {l["title"]}');p('换你来：'+l['followanswer'])
   for n,a in enumerate(l['ans'],1):p(f'{n}  {a}')
   p('带学提示：'+l['reminder'])
 page('参考资料与使用说明');p('本书借鉴公开资料中故事导入、图解、跟练和游戏的呈现方式，按所提供的七讲大纲独立编写。人物、情境、题目和插图均为本书原创，不复制原书页面。')
 for n,r in enumerate(data['references'],1):
  h(f'{n}  {r["title"]}');p(r['type']);p('本书参考：'+r['use']);p(r['url'])
 p('资料访问日期：2026 年 10 月 28 日。公开样例并非所提供冬季大纲的完整教材，不代表与其原版难度和授课顺序完全一致。')
 result=out/'一年级逻辑思维森林探险教材.docx';doc.core_properties.author='';doc.core_properties.title=data['title'];doc.save(result)
 # Keep a readable source in Git; drawings and rendered outputs remain ignored.
 md=['# '+data['title'],'\n原创图解故事版，非学而思官方教材。每讲 1 道跟练与 5 道分层练习，共 42 道。']
 for i,l in enumerate(data['lessons'],1):
  md += [f'\n## 第{i}讲 {l["title"]}',f'\n### 故事 {l["scene"]}',l['dialogue'],f'\n任务：{l["mission"]}',f'\n目标：{l["goal"]}',f'\n材料：{l["materials"]}','\n### 讲解',l['explain'],'\n### 思考路线']+[f'{n}. {s}' for n,s in enumerate(l['steps'],1)]+['\n### 示例',l['example'],'\n### 换你来',l['follow'],'\n### 分层练习']+[f'{n}. 【{a}】{q}' for n,(a,q) in enumerate(l['qs'],1)]+['\n### 亲子游戏',l['game'],'\n### 答案与带学提示','换你来：'+l['followanswer']]+[f'{n}. {a}' for n,a in enumerate(l['ans'],1)]+[l['parent'],l['reminder']]
 md+=['\n## 网络参考与范围说明','\n访问日期：2026-10-28。实际查看公开一年级教程截图、官方课程介绍及旧版公开测试；未读取冬季课程整套教材。']
 for r in data['references']:md += [f'\n- [{r["title"]}]({r["url"]})\n  {r["type"]} {r["use"]}']
 md+=['\n所有正文题目和图示均为独立编写与绘制；不复刻参考书的角色、原题、插图或页面。']
 (ROOT/'docs/grade1-logic-parent-child-workbook.md').write_text('\n'.join(md)+'\n')
 print(result)

if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('--font',type=Path,default=Path('/Library/Fonts/Arial Unicode.ttf'));parser.add_argument('--output',type=Path,default=ROOT/'outputs/grade1-logic-story');args=parser.parse_args();build(args.font,args.output)
